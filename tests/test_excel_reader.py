"""ExcelReader 单元测试。

验证：
- 不遗漏：原始文件中有 8 个参数，读取后全部存在
- 不错误：读取数值与原始值一致
- 多批次：多个批次均正确解析
- 合并单元格场景：批号仅首行出现，后续行为空
- 参数名不匹配的行被忽略
- 空文件不崩溃
"""
import os
import tempfile
import pytest

from src.readers.excel_reader import ExcelReader
from tests.test_helpers import (
    build_excel_file,
    make_standard_header,
    make_batch_rows,
    ALL_8_PARAMS,
)


class TestExcelReaderSingleBatch:
    """单批次基本读取测试。"""

    def _create_single_batch_file(self, params, include_batch_no_only_first=False):
        header = make_standard_header()
        rows = make_batch_rows("B001", "M001", params, include_batch_no_only_first)
        return build_excel_file(rows, header)

    def test_reads_all_8_params_not_missing(self):
        """不遗漏：8 个参数必须全部读取到。"""
        path = self._create_single_batch_file(ALL_8_PARAMS)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert len(batches) == 1
        p = batches[0].process_param
        assert p is not None
        assert p.is_complete(), f"缺失参数: {p.missing_fields()}"

    def test_reads_all_8_params_not_wrong(self):
        """不错误：读取数值必须与原始值完全一致。"""
        path = self._create_single_batch_file(ALL_8_PARAMS)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        p = batches[0].process_param
        assert p.cable_core_od == 10.5
        assert p.sheath_od == 14.0
        assert p.extrusion_inner == 9.0
        assert p.extrusion_outer == 13.0
        assert p.screw_speed == 120.0
        assert p.screw_current == 45.0
        assert p.prod_speed == 60.0
        assert p.actual_prod_speed == 58.5

    def test_batch_no_and_material_no_correct(self):
        """批号和物料品号正确解析。"""
        path = self._create_single_batch_file(ALL_8_PARAMS)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert batches[0].batch_no == "B001"
        assert batches[0].material_no == "M001"

    def test_source_file_recorded(self):
        """来源文件名被记录在 process_param 中。"""
        path = self._create_single_batch_file(ALL_8_PARAMS)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert batches[0].process_param.source_file == os.path.basename(path)

    def test_merged_cell_style_batch_no_only_first_row(self):
        """合并单元格场景：批号仅首行出现，后续行为 None，仍能正确聚合。"""
        path = self._create_single_batch_file(ALL_8_PARAMS, include_batch_no_only_first=True)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert len(batches) == 1
        assert batches[0].process_param.is_complete()

    def test_integer_values_parsed_as_float(self):
        """整数值（如 120）应被解析为 float。"""
        params = {**ALL_8_PARAMS, "螺杆速度": 120}
        path = self._create_single_batch_file(params)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert batches[0].process_param.screw_speed == 120.0


class TestExcelReaderMultipleBatches:
    """多批次读取测试：不遗漏、不串批。"""

    def _create_two_batch_file(self):
        header = make_standard_header()
        params_b1 = {**ALL_8_PARAMS, "缆芯外径": 10.0, "护套外径": 13.0}
        params_b2 = {**ALL_8_PARAMS, "缆芯外径": 11.0, "护套外径": 15.0}
        rows = (
            make_batch_rows("B001", "M001", params_b1) +
            make_batch_rows("B002", "M001", params_b2)
        )
        return build_excel_file(rows, header), params_b1, params_b2

    def test_reads_correct_batch_count(self):
        path, _, _ = self._create_two_batch_file()
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert len(batches) == 2

    def test_no_batch_missing(self):
        """两个批次都不遗漏。"""
        path, _, _ = self._create_two_batch_file()
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        batch_nos = {b.batch_no for b in batches}
        assert "B001" in batch_nos
        assert "B002" in batch_nos

    def test_no_param_cross_contamination(self):
        """两批次参数不能串批：B001 和 B002 的缆芯外径各自独立。"""
        path, params_b1, params_b2 = self._create_two_batch_file()
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        by_no = {b.batch_no: b for b in batches}
        assert by_no["B001"].process_param.cable_core_od == params_b1["缆芯外径"]
        assert by_no["B002"].process_param.cable_core_od == params_b2["缆芯外径"]

    def test_all_8_params_per_batch(self):
        """每个批次各自有 8 个完整参数。"""
        path, _, _ = self._create_two_batch_file()
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        for batch in batches:
            assert batch.process_param.is_complete(), (
                f"批号 {batch.batch_no} 缺失参数: {batch.process_param.missing_fields()}"
            )

    def test_multiple_materials(self):
        """不同物料品号的批次能正确区分。"""
        header = make_standard_header()
        rows = (
            make_batch_rows("B001", "M001", ALL_8_PARAMS) +
            make_batch_rows("B002", "M002", ALL_8_PARAMS)
        )
        path = build_excel_file(rows, header)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        by_no = {b.batch_no: b for b in batches}
        assert by_no["B001"].material_no == "M001"
        assert by_no["B002"].material_no == "M002"


class TestExcelReaderEdgeCases:
    """边界场景测试。"""

    def test_unknown_param_name_ignored(self):
        """不在 8 个标准参数中的行应被忽略。"""
        header = make_standard_header()
        rows = make_batch_rows("B001", "M001", ALL_8_PARAMS)
        rows.append(("B001", "M001", "未知参数X", 99.9))
        path = build_excel_file(rows, header)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert len(batches) == 1
        # 未知参数不会影响 8 个标准参数的完整性
        assert batches[0].process_param.is_complete()

    def test_partial_params_batch_still_returned(self):
        """只有部分参数的批次仍然返回（清洗在 DataCleaner 中做，Reader 不过滤）。"""
        header = make_standard_header()
        partial = {"缆芯外径": 10.5, "护套外径": 14.0}
        rows = make_batch_rows("B001", "M001", partial)
        path = build_excel_file(rows, header)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert len(batches) == 1
        assert not batches[0].process_param.is_complete()

    def test_empty_sheet_returns_empty_list(self):
        """空工作表不崩溃，返回空列表。"""
        import openpyxl, tempfile, os
        tmp = tempfile.mkdtemp()
        path = os.path.join(tmp, "empty.xlsx")
        wb = openpyxl.Workbook()
        wb.save(path)
        reader = ExcelReader(tmp, "empty.xlsx")
        batches = reader.read_file(path)
        assert batches == []

    def test_no_files_in_dir_returns_empty_list(self):
        """目录中无匹配文件，返回空列表，不崩溃。"""
        tmp = tempfile.mkdtemp()
        reader = ExcelReader(tmp, "HT_工序*.xlsx")
        batches = reader.read_all()
        assert batches == []

    def test_non_numeric_param_value_becomes_none(self):
        """非数字参数值（如空格、横线）解析为 None。"""
        header = make_standard_header()
        rows = [("B001", "M001", "缆芯外径", "N/A")]
        path = build_excel_file(rows, header)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert batches[0].process_param.cable_core_od is None

    def test_float_string_parsed_correctly(self):
        """Excel 中以字符串形式存储的浮点数能正确解析。"""
        header = make_standard_header()
        rows = make_batch_rows("B001", "M001", {**ALL_8_PARAMS})
        # 替换一个值为字符串
        rows[0] = (rows[0][0], rows[0][1], rows[0][2], "10.5")
        path = build_excel_file(rows, header)
        reader = ExcelReader(os.path.dirname(path), os.path.basename(path))
        batches = reader.read_file(path)
        assert batches[0].process_param.cable_core_od == 10.5

    def test_read_all_aggregates_multiple_files(self):
        """read_all 能聚合多个文件的批次，不遗漏。"""
        import shutil
        tmp = tempfile.mkdtemp()
        header = make_standard_header()

        params1 = {**ALL_8_PARAMS, "缆芯外径": 10.0}
        rows1 = make_batch_rows("B001", "M001", params1)
        path1 = os.path.join(tmp, "HT_工序2024_01.xlsx")
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(header)
        for r in rows1:
            ws.append(r)
        wb.save(path1)

        params2 = {**ALL_8_PARAMS, "缆芯外径": 11.0}
        rows2 = make_batch_rows("B002", "M001", params2)
        path2 = os.path.join(tmp, "HT_工序2024_02.xlsx")
        wb2 = openpyxl.Workbook()
        ws2 = wb2.active
        ws2.append(header)
        for r in rows2:
            ws2.append(r)
        wb2.save(path2)

        reader = ExcelReader(tmp, "HT_工序*.xlsx")
        batches = reader.read_all()
        assert len(batches) == 2
        batch_nos = {b.batch_no for b in batches}
        assert "B001" in batch_nos
        assert "B002" in batch_nos
